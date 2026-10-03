"""Expand the GH1 NIMS explicit-negative pool and quantify repeated units.

Only source cells literally recorded as ``<0.1`` are included. The primary
pool is restricted to Table 3 Expression_binary == 1; all other explicit
below-threshold cells are counted separately, not silently discarded.
"""

from __future__ import annotations

import csv
import math
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
SOURCE_DIR = ROOT / "research" / "sources" / "gh1-2014"
NIMS_FILE = SOURCE_DIR / "cb500244v_si_003.xlsx"
ANNOTATION_FILE = SOURCE_DIR / "cb500244v_si_004.xlsx"
SAMPLE_FILE = HERE / "gh1_nims_negative_sample_100.csv"
POOL_FILE = HERE / "gh1_nims_candidate_pool_1803.csv"
ALL_NEGATIVES_FILE = HERE / "gh1_nims_negative_candidates_all_2076.csv"
BINARY0_FILE = HERE / "gh1_nims_expression_binary0_candidates_273.csv"
PAIR_FILE = HERE / "gh1_nims_candidate_pairs.csv"
SUMMARY_FILE = HERE / "gh1_nims_candidate_pool_summary.csv"
CONDITION_FILE = HERE / "gh1_nims_source_condition_key_audit.csv"
PAIR_CONTEXT_FILE = HERE / "gh1_nims_pair_context_audit.csv"
SUBSTRATE_COLUMNS = ((3, 6, "xylobiose"), (4, 7, "cellobiose"), (5, 8, "lactose"))
VALID_PH = {5, 8}
VALID_TEMP = {40, 60, 80, 90}
ACCESSION_PATTERN = re.compile(r"^([A-Z]{1,4}\d+\.\d+)_")


def accession_from_protein_id(protein_id: str) -> str:
    """Return a versioned GenBank accession only when the label has one."""
    match = ACCESSION_PATTERN.match(protein_id)
    return match.group(1) if match else ""


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    if not rows:
        raise ValueError(f"Refusing to write empty table: {path}")
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        fields = list(dict.fromkeys(key for row in rows for key in row))
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    annotation_sheet = load_workbook(ANNOTATION_FILE, read_only=True, data_only=True).active
    annotations = {
        row[0]: {
            "kingdom": row[1],
            "soluble_expression_mg_ml": row[2],
            "expression_binary": row[3],
        }
        for row in annotation_sheet.iter_rows(min_row=7, values_only=True)
        if row[0] is not None
    }

    sheet = load_workbook(NIMS_FILE, read_only=True, data_only=True).active
    pool: list[dict[str, object]] = []
    all_negative_cells: list[dict[str, object]] = []
    all_counts: Counter[str] = Counter()
    all_proteins: dict[str, set[str]] = defaultdict(set)
    all_pairs: dict[str, set[tuple[str, str]]] = defaultdict(set)
    condition_rows_by_protein: dict[str, list[tuple[int, int, int]]] = defaultdict(list)
    negatives_by_protein: Counter[str] = Counter()
    pair_outcomes: dict[tuple[str, str], Counter[str]] = defaultdict(Counter)
    pair_conditions: dict[tuple[str, str], set[tuple[int, int]]] = defaultdict(set)
    current_protein = None
    for excel_row, row in enumerate(sheet.iter_rows(min_row=7, values_only=True), start=7):
        if row[0] is not None:
            current_protein = row[0]
        if current_protein is None or row[1] not in VALID_PH or row[2] not in VALID_TEMP:
            continue
        condition_rows_by_protein[str(current_protein)].append(
            (int(row[1]), int(row[2]), excel_row)
        )
        annotation = annotations.get(current_protein)
        status = (
            "expression_binary=1" if annotation and annotation["expression_binary"] == 1
            else "expression_binary=0" if annotation and annotation["expression_binary"] == 0
            else "expression_status_unjoined_or_unknown"
        )

        for value_col, cv_col, substrate in SUBSTRATE_COLUMNS:
            raw = row[value_col]
            if status == "expression_binary=1":
                pair_key = (str(current_protein), substrate)
                pair_conditions[pair_key].add((int(row[1]), int(row[2])))
                if isinstance(raw, str) and raw.strip() == "<0.1":
                    outcome = "explicit_below_background"
                elif isinstance(raw, (int, float)) and math.isclose(float(raw), 0.1, abs_tol=1e-12):
                    outcome = "exact_threshold_boundary"
                elif isinstance(raw, (int, float)) and float(raw) > 0.1:
                    outcome = "numeric_above_background"
                elif isinstance(raw, (int, float)):
                    outcome = "numeric_below_background_not_censored_string"
                else:
                    outcome = "missing_or_non_numeric"
                pair_outcomes[pair_key][outcome] += 1
            if not (isinstance(raw, str) and raw.strip() == "<0.1"):
                continue
            all_counts[status] += 1
            all_proteins[status].add(str(current_protein))
            all_pairs[status].add((str(current_protein), substrate))
            negatives_by_protein[str(current_protein)] += 1
            record = {
                    "protein_accession_species": current_protein,
                    "ncbi_accession": accession_from_protein_id(str(current_protein)),
                    "substrate": substrate,
                    "assay_modality": "NIMS",
                    "pH": int(row[1]),
                    "temperature_C": int(row[2]),
                    "mean_conversion_rate_raw": raw.strip(),
                    "label_interpretation": "below reported 0.1 conversion background; not numeric zero",
                    "background_threshold_conversion_rate": 0.1,
                    "biological_replicates_reported": 3,
                    "coefficient_of_variation_raw": row[cv_col],
                    "kingdom": annotation["kingdom"] if annotation else "",
                    "soluble_expression_mg_ml": annotation["soluble_expression_mg_ml"] if annotation else "",
                    "expression_binary": annotation["expression_binary"] if annotation else "",
                    "expression_status": status,
                    "included_in_core_candidate_pool": str(status == "expression_binary=1").lower(),
                    "source_article_doi": "10.1021/cb500244v",
                    "source_supplement_doi": "10.1021/cb500244v.s003",
                    "source_file": NIMS_FILE.name,
                    "source_sheet": sheet.title,
                    "source_excel_row": excel_row,
                    "source_excel_column": {3: "D", 4: "E", 5: "F"}[value_col],
                    "source_cell": f"{{col}}{{row}}".format(
                        col={3: "D", 4: "E", 5: "F"}[value_col], row=excel_row
                    ),
                }
            all_negative_cells.append(record)
            if status == "expression_binary=1":
                pool.append(record)

    if len(pool) != 1803:
        raise RuntimeError(f"Expected 1,803 expression-positive <0.1 cells, got {len(pool)}")
    if len(all_negative_cells) != 2076:
        raise RuntimeError(f"Expected 2,076 total explicit <0.1 cells, got {len(all_negative_cells)}")
    observation_keys = [
        (r["protein_accession_species"], r["substrate"], r["pH"], r["temperature_C"])
        for r in pool
    ]
    condition_key_counts = Counter(observation_keys)
    duplicate_condition_key_groups = {
        key: count for key, count in condition_key_counts.items() if count > 1
    }

    with SAMPLE_FILE.open(encoding="utf-8-sig", newline="") as stream:
        sample = list(csv.DictReader(stream))
    pool_key_set = set(observation_keys)
    sample_missing = [
        row["pilot_id"]
        for row in sample
        if (
            row["protein_accession_species"], row["substrate"],
            int(row["pH"]), int(row["temperature_C"])
        ) not in pool_key_set
    ]
    if sample_missing:
        raise RuntimeError(f"Existing 100-row sample not contained in pool: {sample_missing[:5]}")

    pair_groups: dict[tuple[str, str], list[dict[str, object]]] = defaultdict(list)
    for row in pool:
        pair_groups[(str(row["protein_accession_species"]), str(row["substrate"]))].append(row)
    pair_rows = []
    for (protein, substrate), rows in sorted(pair_groups.items()):
        conditions = sorted({(int(r["pH"]), int(r["temperature_C"])) for r in rows})
        pair_rows.append(
            {
                "protein_accession_species": protein,
                "ncbi_accession": accession_from_protein_id(protein),
                "substrate": substrate,
                "explicit_below_background_condition_cells": len(rows),
                "distinct_pH_temperature_conditions": len(conditions),
                "conditions_pH_C": ";".join(f"{ph}/{temp}" for ph, temp in conditions),
                "source_article_doi": "10.1021/cb500244v",
                "source_supplement_doi": "10.1021/cb500244v.s003",
            }
        )

    pair_context_rows = []
    for (protein, substrate), outcomes in sorted(pair_outcomes.items()):
        neg = outcomes["explicit_below_background"]
        above = outcomes["numeric_above_background"]
        boundary = outcomes["exact_threshold_boundary"]
        if neg and above:
            pair_state = "mixed_below_and_above_across_conditions"
        elif neg and not above:
            pair_state = "explicit_below_background_no_above_readout"
        elif above and not neg:
            pair_state = "above_background_only"
        else:
            pair_state = "boundary_or_missing_only"
        pair_context_rows.append(
            {
                "protein_accession_species": protein,
                "ncbi_accession": accession_from_protein_id(protein),
                "substrate": substrate,
                "distinct_pH_temperature_conditions": len(pair_conditions[(protein, substrate)]),
                "raw_condition_cells_including_source_duplicates": sum(outcomes.values()),
                "explicit_below_background_cells": neg,
                "numeric_above_background_cells": above,
                "exact_threshold_boundary_cells": boundary,
                "numeric_below_background_not_censored_string_cells": outcomes["numeric_below_background_not_censored_string"],
                "missing_or_non_numeric_cells": outcomes["missing_or_non_numeric"],
                "pair_context_state": pair_state,
            }
        )

    expected_conditions = {(ph, temp) for ph in VALID_PH for temp in VALID_TEMP}
    source_condition_rows = []
    for protein, rows in sorted(condition_rows_by_protein.items()):
        counts = Counter((ph, temp) for ph, temp, _ in rows)
        missing = sorted(expected_conditions - set(counts))
        duplicate = sorted((key, count) for key, count in counts.items() if count > 1)
        annotation = annotations.get(protein, {})
        source_condition_rows.append(
            {
                "protein_accession_species": protein,
                "ncbi_accession": accession_from_protein_id(protein),
                "expression_binary": annotation.get("expression_binary", ""),
                "source_condition_rows": len(rows),
                "distinct_pH_temperature_keys": len(counts),
                "expected_unique_keys": len(expected_conditions),
                "missing_pH_temperature_keys": ";".join(f"{ph}/{temp}" for ph, temp in missing),
                "duplicate_pH_temperature_keys": ";".join(
                    f"{ph}/{temp}×{count}" for (ph, temp), count in duplicate
                ),
                "explicit_below_background_cells": negatives_by_protein[protein],
                "source_rows": ";".join(str(excel_row) for _, _, excel_row in rows),
            }
        )

    substrate_rows = []
    for _, _, substrate in SUBSTRATE_COLUMNS:
        rows = [r for r in pool if r["substrate"] == substrate]
        pair_subset = [r for r in pair_rows if r["substrate"] == substrate]
        substrate_rows.append(
            {
                "expression_status": "expression_binary=1",
                "substrate": substrate,
                "explicit_below_background_cells": len(rows),
                "proteins_with_at_least_one_negative_cell": len({r["protein_accession_species"] for r in rows}),
                "negative_enzyme_substrate_pairs": len(pair_subset),
                "negative_pair_condition_cells": sum(int(r["distinct_pH_temperature_conditions"]) for r in pair_subset),
            }
        )

    total_unique_proteins = len({r["protein_accession_species"] for r in pool})
    total_unique_pairs = len(pair_rows)
    core_cells = all_counts["expression_binary=1"]
    summary_rows = [
        {"metric": "explicit_<0.1_cells", "count": sum(all_counts.values()), "definition": "All Table 2 cells literally equal to <0.1 across joined expression-status groups"},
        {"metric": "expression_binary_1_cells", "count": core_cells, "definition": "Candidate pool retained by the existing conservative expression-status filter"},
        {"metric": "expression_binary_0_cells", "count": all_counts["expression_binary=0"], "definition": "Measured below-background cells retained as a separate low-expression stratum; see audit_expression_binary0_context.py"},
        {"metric": "unjoined_or_unknown_status_cells", "count": all_counts["expression_status_unjoined_or_unknown"], "definition": "Explicit <0.1 cells without a joined binary expression status"},
        {"metric": "expression_binary_1_unique_proteins", "count": total_unique_proteins, "definition": "Distinct literal Table 2 protein accession/species identifiers with at least one candidate negative cell"},
        {"metric": "expression_binary_1_negative_enzyme_substrate_pairs", "count": total_unique_pairs, "definition": "Distinct protein accession/species × substrate pairs with ≥1 explicit <0.1 condition cell"},
        {"metric": "expression_binary_1_unique_protein_substrate_pH_temperature_keys", "count": len(condition_key_counts), "definition": "Distinct condition keys; may be lower than source-cell rows if the supplement repeats a protein-condition key"},
        {"metric": "duplicate_protein_substrate_pH_temperature_key_groups", "count": len(duplicate_condition_key_groups), "definition": "Number of condition keys represented by multiple explicit <0.1 source cells"},
        {"metric": "extra_source_cells_on_duplicate_condition_keys", "count": sum(n - 1 for n in duplicate_condition_key_groups.values()), "definition": "Additional source cells beyond one per duplicated condition key; retained, not deduplicated"},
        {"metric": "condition_cells_per_negative_pair_mean", "count": round(core_cells / total_unique_pairs, 3), "definition": "Mean pH-temperature cells below background per negative enzyme-substrate pair; not independent proteins"},
        {"metric": "reported_biological_replicates_per_cell", "count": 3, "definition": "Triplicate design reported by paper; individual replicate values are not present in this supplement table"},
        {"metric": "nominal_replicate_measurements_for_expression_binary_1_cells", "count": core_cells * 3, "definition": "Arithmetic only; cannot call all replicates individually negative because only cell-level mean/censored result is supplied"},
        {"metric": "sampled_rows_present_in_candidate_pool", "count": len(sample) - len(sample_missing), "definition": "Existing 100-row sample checked by exact protein-substrate-pH-temperature key"},
    ]
    pair_state_counts = Counter(row["pair_context_state"] for row in pair_context_rows)
    summary_rows.extend(
        {
            "metric": f"enzyme_substrate_pairs__{state}",
            "count": count,
            "definition": "Pair-level outcome class across all Table 2 pH-temperature contexts; not an independent-paper/generalization count",
        }
        for state, count in sorted(pair_state_counts.items())
    )

    write_csv(POOL_FILE, pool)
    write_csv(ALL_NEGATIVES_FILE, all_negative_cells)
    write_csv(BINARY0_FILE, [r for r in all_negative_cells if r["expression_status"] == "expression_binary=0"])
    write_csv(PAIR_FILE, pair_rows)
    write_csv(SUMMARY_FILE, summary_rows + substrate_rows)
    write_csv(CONDITION_FILE, source_condition_rows)
    write_csv(PAIR_CONTEXT_FILE, pair_context_rows)
    print(f"explicit_below_background_cells={sum(all_counts.values())}")
    print(f"by_expression_status={dict(all_counts)}")
    print(f"candidate_pool_cells={core_cells}")
    print(f"candidate_unique_proteins={total_unique_proteins}")
    print(f"candidate_negative_enzyme_substrate_pairs={total_unique_pairs}")
    print(f"unique_condition_keys={len(condition_key_counts)}")
    print(f"duplicate_condition_key_groups={duplicate_condition_key_groups}")
    print(f"source_ids={len(source_condition_rows)}")
    print(f"source_ids_with_all_8_unique_conditions={sum(int(r['distinct_pH_temperature_keys']) == 8 for r in source_condition_rows)}")
    print(f"source_ids_with_duplicate_condition_rows={sum(bool(r['duplicate_pH_temperature_keys']) for r in source_condition_rows)}")
    print(f"pair_context_states={dict(pair_state_counts)}")
    print(f"sample_membership={len(sample)-len(sample_missing)}/{len(sample)}")
    print(f"outputs={ALL_NEGATIVES_FILE}; {POOL_FILE}; {BINARY0_FILE}; {PAIR_FILE}; {PAIR_CONTEXT_FILE}; {SUMMARY_FILE}; {CONDITION_FILE}")


if __name__ == "__main__":
    main()
