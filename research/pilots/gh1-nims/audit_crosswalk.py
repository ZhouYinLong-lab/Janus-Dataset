"""Reconcile GH1 IDs and expression/activity fields across SI Tables 1–3."""

from __future__ import annotations

import csv
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[3]
SOURCE_DIR = ROOT / "research" / "sources" / "gh1-2014"
TABLE1 = SOURCE_DIR / "cb500244v_si_002.xlsx"
TABLE2 = SOURCE_DIR / "cb500244v_si_003.xlsx"
TABLE3 = SOURCE_DIR / "cb500244v_si_004.xlsx"
OUTPUT = Path(__file__).resolve().parent / "gh1_nims_crosswalk_audit.csv"


def main() -> None:
    table1 = load_workbook(TABLE1, data_only=True, read_only=True).active
    table2 = load_workbook(TABLE2, data_only=True, read_only=True).active
    table3 = load_workbook(TABLE3, data_only=True, read_only=True).active

    candidates = {
        row[1]: {"kingdom": row[2], "soluble_expression_mg_ml": row[4]}
        for row in table1.iter_rows(min_row=5, values_only=True)
        if row[1] is not None
    }
    annotations = {
        row[0]: {
            "expression_binary": row[3],
            "nims_max_temp": row[4],
            "cellobiose_nims_activity_binary": row[5],
            "pnp_a400": row[6],
        }
        for row in table3.iter_rows(min_row=7, values_only=True)
        if row[0] is not None
    }
    conditions: dict[str, list[tuple[object, ...]]] = defaultdict(list)
    protein_id = None
    for excel_row, row in enumerate(table2.iter_rows(min_row=7, values_only=True), start=7):
        if row[0] is not None:
            protein_id = row[0]
        if protein_id is not None and row[1] in (5, 8) and row[2] in (40, 60, 80, 90):
            conditions[protein_id].append((excel_row, *row[1:9]))

    if set(candidates) != set(annotations):
        raise RuntimeError("Table 1 and Table 3 protein identifiers do not match exactly")

    records = []
    for protein_id in sorted(candidates):
        candidate = candidates[protein_id]
        annotation = annotations[protein_id]
        rows = conditions.get(protein_id, [])
        explicit_below_background = sum(
            isinstance(row[column], str) and row[column].strip() == "<0.1"
            for row in rows
            for column in (3, 4, 5)
        )
        records.append({
            "protein_accession_species": protein_id,
            "kingdom": candidate["kingdom"],
            "soluble_expression_mg_ml": candidate["soluble_expression_mg_ml"],
            "expression_binary": annotation["expression_binary"],
            "nims_max_temp": annotation["nims_max_temp"],
            "cellobiose_nims_activity_binary": annotation["cellobiose_nims_activity_binary"],
            "pnp_a400": annotation["pnp_a400"],
            "table2_condition_rows": len(rows),
            "table2_explicit_lt_0_1_cells": explicit_below_background,
            "table2_present": bool(rows),
        })

    with OUTPUT.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(records[0]))
        writer.writeheader()
        writer.writerows(records)

    binary_counts = Counter(record["expression_binary"] for record in records)
    concentration_ranges = {
        flag: [record["soluble_expression_mg_ml"] for record in records
               if record["expression_binary"] == flag]
        for flag in (0, 1)
    }
    present = [record for record in records if record["table2_present"]]
    absent_positive = [record for record in records
                       if record["expression_binary"] == 1 and not record["table2_present"]]
    present_zero = [record for record in present if record["expression_binary"] == 0]

    print(f"Table 1/Table 3 exact ID matches: {len(records)}")
    print(f"Table 2 proteins/condition rows: {len(conditions)}/{sum(map(len, conditions.values()))}")
    print(f"Expression_binary counts: {dict(binary_counts)}")
    for flag, values in concentration_ranges.items():
        print(f"Expression_binary={flag} concentration range: {min(values)}–{max(values)} mg/ml")
    print(f"Table 2 expression_binary=0 proteins: {len(present_zero)}")
    print(f"Expression_binary=1 proteins absent from Table 2: {len(absent_positive)}")
    print(f"Explicit <0.1 cells in Table 2 across all represented proteins: "
          f"{sum(record['table2_explicit_lt_0_1_cells'] for record in present)}")
    print(f"Crosswalk CSV: {OUTPUT}")


if __name__ == "__main__":
    main()
