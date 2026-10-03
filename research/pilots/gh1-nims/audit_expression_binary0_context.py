"""Audit measured GH1 NIMS outcomes for proteins flagged Expression_binary=0."""

from __future__ import annotations

import csv
import math
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
SOURCE_DIR = ROOT / "research" / "sources" / "gh1-2014"
NIMS_FILE = SOURCE_DIR / "cb500244v_si_003.xlsx"
ANNOTATION_FILE = SOURCE_DIR / "cb500244v_si_004.xlsx"
OUTPUT = HERE / "gh1_nims_expression_binary0_pair_context.csv"
SUBSTRATES = ((3, "xylobiose"), (4, "cellobiose"), (5, "lactose"))
VALID_PH = {5, 8}
VALID_TEMP = {40, 60, 80, 90}
ACCESSION_PATTERN = re.compile(r"^([A-Z]{1,4}\d+\.\d+)_")


def accession_from_protein_id(protein_id: str) -> str:
    match = ACCESSION_PATTERN.match(protein_id)
    return match.group(1) if match else ""


def outcome(value: object) -> str:
    if isinstance(value, str) and value.strip() == "<0.1":
        return "explicit_below_background"
    if isinstance(value, (int, float)) and math.isclose(float(value), 0.1, abs_tol=1e-12):
        return "exact_threshold_boundary"
    if isinstance(value, (int, float)) and float(value) > 0.1:
        return "numeric_above_background"
    if isinstance(value, (int, float)):
        return "numeric_below_background"
    return "missing_or_non_numeric"


def main() -> None:
    annotation_sheet = load_workbook(
        ANNOTATION_FILE, read_only=True, data_only=True
    ).active
    annotations = {
        row[0]: {
            "soluble_expression_mg_ml": row[2],
            "expression_binary": row[3],
            "cellobiose_nims_activity_binary": row[5],
        }
        for row in annotation_sheet.iter_rows(min_row=7, values_only=True)
        if row[0] is not None
    }

    sheet = load_workbook(NIMS_FILE, read_only=True, data_only=True).active
    pair_data: dict[tuple[str, str], dict[str, object]] = {}
    proteins: set[str] = set()
    current_protein = None
    for excel_row, row in enumerate(sheet.iter_rows(min_row=7, values_only=True), start=7):
        if row[0] is not None:
            current_protein = row[0]
        if current_protein is None or row[1] not in VALID_PH or row[2] not in VALID_TEMP:
            continue
        annotation = annotations.get(current_protein)
        if not annotation or annotation["expression_binary"] != 0:
            continue
        protein = str(current_protein)
        proteins.add(protein)
        for col, substrate in SUBSTRATES:
            key = (protein, substrate)
            record = pair_data.setdefault(
                key,
                {
                    "counts": Counter(),
                    "conditions": set(),
                    "source_rows": [],
                    "negative_cells": [],
                },
            )
            record["conditions"].add((int(row[1]), int(row[2])))
            record["source_rows"].append(excel_row)
            label = outcome(row[col])
            record["counts"][label] += 1
            if label == "explicit_below_background":
                record["negative_cells"].append((excel_row, row[1], row[2], col))

    rows: list[dict[str, object]] = []
    total_explicit = 0
    for (protein, substrate), record in sorted(pair_data.items()):
        annotation = annotations[protein]
        counts: Counter[str] = record["counts"]
        explicit = counts["explicit_below_background"]
        total_explicit += explicit
        if explicit and counts["numeric_above_background"]:
            state = "mixed_below_and_above_across_conditions"
        elif explicit:
            state = "below_background_observed_no_above_in_table"
        elif counts["numeric_above_background"]:
            state = "above_background_only"
        else:
            state = "boundary_or_missing_only"
        rows.append(
            {
                "protein_accession_species": protein,
                "ncbi_accession": accession_from_protein_id(protein),
                "substrate": substrate,
                "expression_binary": annotation["expression_binary"],
                "soluble_expression_mg_ml": annotation["soluble_expression_mg_ml"],
                "cellobiose_nims_activity_binary": annotation["cellobiose_nims_activity_binary"],
                "condition_rows": len(record["source_rows"]),
                "distinct_pH_temperature_conditions": len(record["conditions"]),
                "explicit_below_background_cells": explicit,
                "numeric_above_background_cells": counts["numeric_above_background"],
                "exact_threshold_boundary_cells": counts["exact_threshold_boundary"],
                "numeric_below_background_cells": counts["numeric_below_background"],
                "missing_or_non_numeric_cells": counts["missing_or_non_numeric"],
                "pair_context_state": state,
                "source_rows": ";".join(map(str, record["source_rows"])),
                "negative_source_cells": ";".join(
                    f"{ {3: 'D', 4: 'E', 5: 'F'}[col] }{excel_row}"
                    for excel_row, _, _, col in record["negative_cells"]
                ),
                "source_supplement_doi": "10.1021/cb500244v.s003",
            }
        )

    if len(proteins) != 12 or total_explicit != 273 or len(rows) != 36:
        raise RuntimeError(
            f"Unexpected binary-zero audit: proteins={len(proteins)}, "
            f"explicit_cells={total_explicit}, pairs={len(rows)}"
        )
    with OUTPUT.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)

    state_counts = Counter(row["pair_context_state"] for row in rows)
    concentration_counts = Counter(
        str(annotations[protein]["soluble_expression_mg_ml"]) for protein in proteins
    )
    print(f"Expression_binary=0 proteins with Table 2 rows: {len(proteins)}")
    print(f"Pairs across three substrates: {len(rows)}")
    print(f"Literal <0.1 source cells: {total_explicit}")
    print(f"Pair-context states: {dict(state_counts)}")
    print(f"Soluble-expression range (mg/ml): {min(map(float, concentration_counts))}–"
          f"{max(map(float, concentration_counts))}")
    print(f"Output: {OUTPUT}")


if __name__ == "__main__":
    main()
