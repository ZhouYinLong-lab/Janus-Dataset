"""Create a deterministic, stratified sample of explicit GH1 NIMS non-detections.

The script does not infer negatives from missing rows or numeric zeroes. It selects
only source cells explicitly recorded as ``<0.1`` and restricts to proteins with
Expression_binary == 1 in the linked Supplementary Table 3.
"""

from __future__ import annotations

import csv
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[3]
SOURCE_DIR = ROOT / "research" / "sources" / "gh1-2014"
NIMS_FILE = SOURCE_DIR / "cb500244v_si_003.xlsx"
ANNOTATION_FILE = SOURCE_DIR / "cb500244v_si_004.xlsx"
OUTPUT_FILE = Path(__file__).resolve().parent / "gh1_nims_negative_sample_100.csv"

SUBSTRATES = ((3, 6, "xylobiose"), (4, 7, "cellobiose"), (5, 8, "lactose"))
CONDITIONS = [(substrate, ph, temp) for substrate in ("xylobiose", "cellobiose", "lactose")
              for ph in (5, 8) for temp in (40, 60, 80, 90)]
# Four extra observations are distributed across substrates and pH values.
EXTRA_QUOTAS = {
    ("xylobiose", 5, 40),
    ("xylobiose", 8, 40),
    ("cellobiose", 5, 40),
    ("lactose", 8, 40),
}


def main() -> None:
    annotation_sheet = load_workbook(ANNOTATION_FILE, read_only=True, data_only=True).active
    annotations = {
        row[0]: {"kingdom": row[1], "soluble_expression_mg_ml": row[2],
                 "expression_binary": row[3]}
        for row in annotation_sheet.iter_rows(min_row=7, values_only=True)
        if row[0] is not None
    }

    nims_sheet = load_workbook(NIMS_FILE, read_only=True, data_only=True).active
    strata: dict[tuple[str, int, int], list[dict[str, object]]] = defaultdict(list)
    protein_id = None
    for excel_row, row in enumerate(nims_sheet.iter_rows(min_row=7, values_only=True), start=7):
        if row[0] is not None:
            protein_id = row[0]
        if protein_id is None or row[1] not in (5, 8) or row[2] not in (40, 60, 80, 90):
            continue
        protein = annotations.get(protein_id)
        if not protein or protein["expression_binary"] != 1:
            continue

        for value_column, cv_column, substrate in SUBSTRATES:
            raw_value = row[value_column]
            if not (isinstance(raw_value, str) and raw_value.strip() == "<0.1"):
                continue
            key = (substrate, int(row[1]), int(row[2]))
            strata[key].append({
                "pilot_id": "",
                "protein_accession_species": protein_id,
                "enzyme_family": "GH1 beta-glucosidase",
                "substrate": substrate,
                "assay_modality": "NIMS",
                "pH": row[1],
                "temperature_C": row[2],
                "mean_conversion_rate_raw": raw_value,
                "label_interpretation": "below reported 0.1 conversion background; do not replace with numeric zero",
                "background_threshold_conversion_rate": "0.1",
                "biological_replicates": 3,
                "coefficient_of_variation_raw": row[cv_column],
                "kingdom": protein["kingdom"],
                "soluble_expression_mg_ml": protein["soluble_expression_mg_ml"],
                "expression_binary": protein["expression_binary"],
                "source_article_doi": "10.1021/cb500244v",
                "source_supplement_doi": "10.1021/cb500244v.s003",
                "source_file": NIMS_FILE.name,
                "source_sheet": nims_sheet.title,
                "source_excel_row": excel_row,
                "source_excel_column": {3: "D", 4: "E", 5: "F"}[value_column],
                "source_cell": f"{{column}}{{row}}".format(
                    column={3: "D", 4: "E", 5: "F"}[value_column], row=excel_row),
            })

    if len(strata) != 24 or any(len(strata[key]) < 4 for key in CONDITIONS):
        raise RuntimeError("Expected at least four explicit negatives in each of 24 strata")

    selected: list[dict[str, object]] = []
    for key in CONDITIONS:
        items = sorted(strata[key], key=lambda x: (str(x["protein_accession_species"]),
                                                    int(x["source_excel_row"])))
        quota = 5 if key in EXTRA_QUOTAS else 4
        # Evenly spaced deterministic sample within each condition/substrate stratum.
        indices = [int((i + 0.5) * len(items) / quota) for i in range(quota)]
        selected.extend(items[index] for index in indices)

    if len(selected) != 100:
        raise RuntimeError(f"Expected 100 sample rows, got {len(selected)}")
    identities = [(r["protein_accession_species"], r["substrate"], r["pH"], r["temperature_C"])
                  for r in selected]
    if len(set(identities)) != 100:
        raise RuntimeError("Duplicate enzyme–substrate–condition observation in sample")
    for index, record in enumerate(selected, start=1):
        record["pilot_id"] = f"GH1NIMS{index:04d}"

    with OUTPUT_FILE.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(selected[0]))
        writer.writeheader()
        writer.writerows(selected)

    print(f"Output: {OUTPUT_FILE}")
    print(f"Candidate explicit <0.1 observations from expressed-protein overlap: "
          f"{sum(map(len, strata.values()))}")
    print(f"Strata sampled: {len(strata)}; sample rows: {len(selected)}")
    print("Sample count by substrate:")
    for substrate in ("xylobiose", "cellobiose", "lactose"):
        print(f"  {substrate}: {sum(r['substrate'] == substrate for r in selected)}")
    print("Sample count by pH/temperature:")
    for key in CONDITIONS:
        print(f"  {key}: {sum((r['substrate'], r['pH'], r['temperature_C']) == key for r in selected)}")


if __name__ == "__main__":
    main()
